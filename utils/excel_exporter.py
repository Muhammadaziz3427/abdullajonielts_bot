"""Excel exporter utility for generating user and bot analytics spreadsheets."""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.db import get_all_users, get_lessons, get_user_progress


async def export_users_to_excel() -> io.BytesIO:
    """Generate an in-memory Excel workbook containing all user details and stats."""
    users = await get_all_users()
    lessons = await get_lessons(active_only=False)
    total_lessons = len(lessons)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Foydalanuvchilar"

    # Header styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    headers = [
        "№",
        "Telegram ID",
        "Foydalanuvchi nomi",
        "To'liq ismi",
        "Obuna holati",
        "Bloklangan",
        "Referral soni",
        "Kim orqali kelgan (ID)",
        "O'zlashtirilgan darslar",
        "Progress (%)",
        "Ro'yxatdan o'tgan sana",
        "Oxirgi faollik",
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data rows
    row_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    for idx, user in enumerate(users, 1):
        user_id = user["user_id"]
        progress_data = await get_user_progress(user_id)
        completed_count = sum(1 for p in progress_data.values() if p.get("is_completed"))
        percent = f"{int(completed_count / total_lessons * 100)}%" if total_lessons > 0 else "0%"

        is_sub_str = "A'zo bo'lgan" if user.get("is_subscribed") else "Obuna yo'q"
        is_ban_str = "Ha (Blokda)" if user.get("is_banned") else "Yo'q"
        username_str = f"@{user['username']}" if user.get("username") else "Mavjud emas"
        referred_by_str = str(user.get("referred_by") or "-")

        created_str = str(user.get("created_at", ""))[:19].replace("T", " ")
        active_str = str(user.get("last_active_at", ""))[:19].replace("T", " ")

        row_values = [
            idx,
            user_id,
            username_str,
            user.get("full_name", ""),
            is_sub_str,
            is_ban_str,
            user.get("referral_count", 0),
            referred_by_str,
            f"{completed_count} / {total_lessons}",
            percent,
            created_str,
            active_str,
        ]

        ws.append(row_values)
        current_row = idx + 1
        ws.row_dimensions[current_row].height = 20

        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = row_font
            cell.border = thin_border
            if col_idx in (1, 2, 5, 6, 7, 8, 9, 10, 11, 12):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
